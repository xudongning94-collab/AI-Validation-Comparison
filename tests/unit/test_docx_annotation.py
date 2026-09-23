from __future__ import annotations

from zipfile import ZipFile

import pytest
from docx import Document

from bid_compare_agent.annotate import DocxAnnotationConfig, annotate_docx
from bid_compare_agent.models.document_ir import SourceLocator
from bid_compare_agent.models.finding import Finding
from bid_compare_agent.parser.docx_parser import parse_docx
from bid_compare_agent.utils.io import file_sha256


def _source_docx(path):
    document = Document()
    document.add_paragraph("第一段需要人工复核。")
    document.add_paragraph("第二段为低风险说明。")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "表格证据"
    document.save(path)


def _finding(
    finding_id: str,
    document_id: str,
    locator: SourceLocator,
    *,
    severity: str = "high",
    finding_type: str = "text_similarity",
) -> Finding:
    return Finding(
        finding_id=finding_id,
        type=finding_type,
        severity=severity,
        document_id=document_id,
        source_locator=locator,
        summary=f"{finding_id} 风险说明",
        score=0.91,
        evidence={"text": "定位证据"},
    )


def test_annotation_writes_native_comments_and_preserves_source(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "annotated.docx"
    _source_docx(source)
    source_ir = parse_docx(source)
    source_hash = file_sha256(source)

    source_finding = _finding(
        "f-source",
        source_ir.document_id,
        SourceLocator(kind="docx_paragraph", paragraph_index=0),
        finding_type="ai_likelihood",
    )
    peer_finding = _finding(
        "f-peer",
        "doc-peer",
        SourceLocator(kind="docx_paragraph", paragraph_index=9),
        severity="critical",
    )
    peer_finding.peer_document_id = source_ir.document_id
    peer_finding.peer_source_locator = SourceLocator(
        kind="docx_table_cell",
        table_index=0,
        row_index=0,
        cell_index=0,
    )
    low_finding = _finding(
        "f-low",
        source_ir.document_id,
        SourceLocator(kind="docx_paragraph", paragraph_index=1),
        severity="info",
    )
    out_of_range = _finding(
        "f-range",
        source_ir.document_id,
        SourceLocator(kind="docx_paragraph", paragraph_index=99),
    )
    other_document = _finding(
        "f-other",
        "doc-other",
        SourceLocator(kind="docx_paragraph", paragraph_index=0),
    )

    result = annotate_docx(
        source,
        output,
        [source_finding, peer_finding, low_finding, out_of_range, other_document],
    )

    assert file_sha256(source) == source_hash
    assert output.exists()
    assert result.summary == {
        "input_findings": 5,
        "annotated_comments": 2,
        "skipped_findings": 3,
        "skip_reasons": {
            "below_min_severity": 1,
            "finding_targets_another_document": 1,
            "paragraph_index_out_of_range": 1,
        },
    }
    assert {item.target_role for item in result.annotations} == {"source", "peer"}
    assert any("必须人工复核" in item.comment_text for item in result.annotations)
    assert result.metadata["source_sha256"] == source_hash
    assert result.metadata["output_sha256"] == file_sha256(output)

    annotated = Document(output)
    assert len(annotated.comments) == 2
    assert {comment.author for comment in annotated.comments} == {"Bid Compare Agent"}
    with ZipFile(output) as package:
        assert "word/comments.xml" in package.namelist()
        document_xml = package.read("word/document.xml")
    assert document_xml.count(b"<w:commentRangeStart") == 2
    assert document_xml.count(b"<w:commentReference") == 2


def test_annotation_deduplicates_and_honors_max_comments(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "annotated.docx"
    _source_docx(source)
    source_ir = parse_docx(source)
    first = _finding(
        "f-first",
        source_ir.document_id,
        SourceLocator(kind="docx_paragraph", paragraph_index=0),
        severity="critical",
    )
    second = _finding(
        "f-second",
        source_ir.document_id,
        SourceLocator(kind="docx_paragraph", paragraph_index=1),
    )

    result = annotate_docx(
        source,
        output,
        [first, first, second],
        DocxAnnotationConfig(max_comments=1),
    )

    assert result.summary["annotated_comments"] == 1
    assert result.summary["skip_reasons"] == {
        "duplicate_finding": 1,
        "max_comments_reached": 1,
    }


def test_annotation_rejects_unsafe_or_invalid_configuration(tmp_path):
    source = tmp_path / "source.docx"
    _source_docx(source)

    with pytest.raises(ValueError, match="输出文件必须与源文件不同"):
        annotate_docx(source, source, [])
    with pytest.raises(ValueError, match="输出必须是 .docx"):
        annotate_docx(source, tmp_path / "result.json", [])
    with pytest.raises(ValueError, match="author 不能为空"):
        DocxAnnotationConfig(author=" ")
