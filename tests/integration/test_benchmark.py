from __future__ import annotations

import io
import json
import subprocess
import sys
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from jsonschema import validate
from PIL import Image

from bid_compare_agent.benchmark import BenchmarkInputError, run_benchmark
from bid_compare_agent.utils.io import file_sha256


ROOT = Path(__file__).resolve().parents[2]


def _image_bytes() -> bytes:
    stream = io.BytesIO()
    Image.new("RGB", (64, 64), color=(20, 80, 160)).save(stream, format="PNG")
    return stream.getvalue()


def _write_docx(path: Path, final_paragraph: str) -> None:
    document = Document()
    document.add_paragraph("项目应建立统一微服务治理流程和持续交付机制。")
    document.add_paragraph("本方案完全响应采购文件要求，并提供实施与质量保障。")
    document.add_paragraph(final_paragraph)
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "共同表格内容"
    document.add_picture(io.BytesIO(_image_bytes()))
    document.save(path)


def _write_pdf(path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        "The project shall establish a unified microservice governance and delivery process.",
    )
    document.save(path)
    document.close()


def _case(tmp_path: Path) -> tuple[dict, dict[str, str]]:
    tender = tmp_path / "tender.pdf"
    old = tmp_path / "old.docx"
    new = tmp_path / "new.docx"
    _write_pdf(tender)
    _write_docx(old, "旧版本包含原始进度安排和风险控制说明。")
    _write_docx(new, "新版本包含更新后的进度安排和风险控制说明。")
    manifest = {
        "benchmark_id": "integration-case",
        "privacy": {"include_raw_text": False},
        "documents": [
            {"id": "tender", "role": "tender_reference", "path": tender.name},
            {"id": "old", "role": "response_old", "path": old.name},
            {"id": "new", "role": "response_new", "path": new.name},
        ],
        "revision_pairs": [
            {
                "id": "old-vs-new",
                "left": "old",
                "right": "new",
                "expectations": {
                    "min_text_repeated_rate": 0.5,
                    "min_image_repeated_rate": 1.0,
                    "min_exact_sequence_ratio": 0.6,
                    "require_file_difference": True,
                },
            }
        ],
        "reference_checks": [
            {"id": "tender-attribution", "reference": "tender", "targets": ["new", "old"]}
        ],
    }
    hashes = {path.name: file_sha256(path) for path in (tender, old, new)}
    return manifest, hashes


def test_benchmark_is_private_schema_valid_and_preserves_sources(tmp_path: Path):
    manifest, hashes = _case(tmp_path)

    result = run_benchmark(manifest, manifest_dir=tmp_path)

    assert result["status"] == "passed"
    assert result["privacy"]["raw_text_included"] is False
    assert result["integrity"]["inputs_unchanged"] is True
    pair = result["revision_pairs"][0]
    assert pair["metrics"]["logical_diff"]["paragraphs"]["changed_blocks"] == 1
    assert pair["metrics"]["logical_diff"]["images"]["perceptual_changed_or_unpaired"] == 0
    assert all(item["passed"] for item in pair["expectation_checks"])
    assert hashes == {
        name: file_sha256(tmp_path / name)
        for name in ("tender.pdf", "old.docx", "new.docx")
    }
    attribution = result["reference_checks"][0]["revision_pair_reference_attribution"]
    assert attribution is not None
    assert attribution["revision_pair_high_pairs"] >= 1
    schema = json.loads((ROOT / "schemas" / "benchmark.schema.json").read_text(encoding="utf-8"))
    validate(result, schema)


def test_benchmark_rejects_raw_text_output(tmp_path: Path):
    manifest, _hashes = _case(tmp_path)
    manifest["privacy"]["include_raw_text"] = True

    with pytest.raises(BenchmarkInputError, match="不允许"):
        run_benchmark(manifest, manifest_dir=tmp_path)


def test_benchmark_cli_writes_report(tmp_path: Path):
    manifest, _hashes = _case(tmp_path)
    manifest_path = tmp_path / "manifest.json"
    report_path = tmp_path / "report.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")

    completed = subprocess.run(
        [
            sys.executable,
            str(ROOT / "scripts" / "run_benchmark.py"),
            str(manifest_path),
            "-o",
            str(report_path),
        ],
        cwd=ROOT,
        capture_output=True,
        text=True,
        check=False,
    )

    assert completed.returncode == 0, completed.stderr
    assert "status=passed" in completed.stdout
    assert json.loads(report_path.read_text(encoding="utf-8"))["benchmark_id"] == "integration-case"
