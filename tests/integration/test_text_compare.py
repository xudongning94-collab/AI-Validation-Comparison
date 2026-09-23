from pathlib import Path

from docx import Document

from bid_compare_agent.compare import TextCompareConfig, compare_document_pair, compare_document_set
from bid_compare_agent.parser.docx_parser import parse_docx


REPEATED = (
    "本项目采用统一能力平台建设模式，通过标准接口接入业务系统，"
    "形成统一身份认证、统一数据管理、统一服务编排和统一运维监控能力。"
)


def _make_doc(path: Path, body: list[str]) -> None:
    doc = Document()
    doc.add_heading("技术方案", level=1)
    for text in body:
        doc.add_paragraph(text)
    doc.save(path)


def test_exact_repeated_paragraph_becomes_high_finding(tmp_path: Path):
    a = tmp_path / "a.docx"
    b = tmp_path / "b.docx"
    _make_doc(a, [REPEATED, "甲方侧建设实施计划采用分阶段推进方式，并按月组织里程碑验收。"])
    _make_doc(b, [REPEATED, "乙方侧提供驻场支持和故障响应机制，保障系统连续稳定运行。"])

    doc_a = parse_docx(a)
    doc_b = parse_docx(b)
    findings, summary = compare_document_pair(
        doc_a,
        doc_b,
        TextCompareConfig(min_chars=15),
    )

    assert summary.high_pairs >= 1
    assert any(f.severity == "high" and f.score == 1.0 for f in findings)
    assert summary.document_a.repeated_rate_high > 0
    assert findings[0].source_locator.paragraph_index is not None
    assert findings[0].peer_source_locator.paragraph_index is not None


def test_three_documents_generate_three_pair_summaries(tmp_path: Path):
    paths = [tmp_path / f"doc{i}.docx" for i in range(3)]
    for i, path in enumerate(paths):
        _make_doc(path, [REPEATED, f"这是第{i}份文件的独立实施说明，内容用于验证三文档两两比对逻辑是否正确。"])
    docs = [parse_docx(path) for path in paths]

    result = compare_document_set(docs, TextCompareConfig(min_chars=15))
    assert len(result.pair_summaries) == 3
    assert all(summary.high_pairs >= 1 for summary in result.pair_summaries)


def test_document_id_is_stable_for_same_file(tmp_path: Path):
    p = tmp_path / "stable.docx"
    _make_doc(p, [REPEATED])
    first = parse_docx(p)
    second = parse_docx(p)
    assert first.document_id == second.document_id
    assert first.document_id.startswith("doc-")


def test_standard_response_is_downweighted_not_high(tmp_path: Path):
    response = "我方承诺完全响应招标文件第十条要求，并按照招标文件规定完成系统建设、验收及服务保障工作。"
    a = tmp_path / "response-a.docx"
    b = tmp_path / "response-b.docx"
    _make_doc(a, [response])
    _make_doc(b, [response])

    findings, summary = compare_document_pair(
        parse_docx(a),
        parse_docx(b),
        TextCompareConfig(min_chars=15, interference_downweight_factor=0.70),
    )

    assert summary.high_pairs == 0
    assert summary.medium_pairs >= 1
    assert findings[0].metadata["raw_similarity"] == 1.0
    assert findings[0].metadata["interference_factor"] == 0.7
