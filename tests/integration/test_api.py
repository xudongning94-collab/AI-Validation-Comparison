from __future__ import annotations

import io
import json
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from jsonschema import validate

from bid_compare_agent.api import create_app


ROOT = Path(__file__).resolve().parents[2]
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _docx_bytes(*paragraphs: str) -> bytes:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _document_files():
    common = "项目实施阶段将建立统一流程和质量保障机制，确保各项任务按计划稳定推进。"
    return [
        ("files", ("a.docx", _docx_bytes(common, "我方完全响应招标要求并提供实施保障。"), DOCX_MIME)),
        ("files", ("b.docx", _docx_bytes(common, "项目团队负责交付和风险跟踪管理。"), DOCX_MIME)),
    ]


def test_health_and_openapi_expose_versioned_routes():
    client = TestClient(create_app())

    health = client.get("/health")
    assert health.status_code == 200
    assert health.json() == {
        "status": "ok",
        "service": "bid-compare-agent",
        "version": "0.1.0-alpha.12",
        "api_version": "1.0.0",
    }

    paths = client.get("/openapi.json").json()["paths"]
    assert {
        "/v1/parse",
        "/v1/preprocess",
        "/v1/text-compare",
        "/v1/image-compare",
        "/v1/analyze",
        "/v1/format-check",
        "/v1/ai-check",
        "/v1/score",
        "/v1/annotate-docx",
    }.issubset(paths)


def test_parse_and_preprocess_endpoints_return_traceable_results():
    client = TestClient(create_app())
    content = _docx_bytes("我方完全响应招标要求并建立项目实施保障机制。")
    upload = {"file": ("source.docx", content, DOCX_MIME)}

    parsed = client.post("/v1/parse", files=upload)
    assert parsed.status_code == 200
    parsed_result = parsed.json()["result"]
    assert parsed_result["filename"] == "source.docx"
    assert parsed_result["document_id"].startswith("doc-")
    assert parsed_result["paragraphs"][0]["source_locator"]["paragraph_index"] == 0

    preprocessed = client.post("/v1/preprocess", files=upload)
    assert preprocessed.status_code == 200
    preprocess_result = preprocessed.json()["result"]
    assert preprocess_result["document"]["document_id"] == parsed_result["document_id"]
    assert preprocess_result["interference"]["summary"]["downweight_count"] == 1


def test_analysis_and_scoring_endpoints_cover_complete_pipeline():
    client = TestClient(create_app())
    expected_types = {
        "/v1/text-compare": "text_similarity",
        "/v1/image-compare": "image_similarity",
        "/v1/format-check": "format_check",
        "/v1/ai-check": "ai_likelihood",
        "/v1/score": "unified_risk",
    }

    for endpoint, result_type in expected_types.items():
        response = client.post(endpoint, files=_document_files())
        assert response.status_code == 200, (endpoint, response.text)
        payload = response.json()
        assert payload["api_version"] == "1.0.0"
        result = payload["result"]
        actual_type = (
            result.get("compare_type")
            or result.get("check_type")
            or result.get("analysis_type")
            or result.get("score_type")
        )
        assert actual_type == result_type


def test_analyze_endpoint_returns_complete_bundle_and_deterministic_report():
    client = TestClient(create_app())

    response = client.post("/v1/analyze", files=_document_files())

    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["operation"] == "analyze"
    result = payload["result"]
    assert result["analysis_type"] == "complete_analysis"
    assert result["metadata"]["parsed_once"] is True
    assert result["report"]["report_id"].startswith("report-")
    assert result["report"]["requires_human_review"] is True
    assert set(result["results"]) == {
        "text_comparison", "image_comparison", "format_checks", "ai_likelihood", "scoring"
    }


def test_annotate_endpoint_returns_native_word_comments():
    client = TestClient(create_app())
    content = _docx_bytes("该段落需要写入 Word 原生批注。")
    parsed = client.post(
        "/v1/parse",
        files={"file": ("source.docx", content, DOCX_MIME)},
    ).json()["result"]
    finding = {
        "finding_id": "f-api-annotation",
        "type": "format_issue",
        "severity": "high",
        "document_id": parsed["document_id"],
        "source_locator": {"kind": "docx_paragraph", "paragraph_index": 0},
        "summary": "API 批注测试",
        "score": 0.9,
        "evidence": {"text": "该段落需要写入 Word 原生批注。"},
    }

    response = client.post(
        "/v1/annotate-docx",
        files={"file": ("source.docx", content, DOCX_MIME)},
        data={"findings_json": json.dumps({"findings": [finding]}, ensure_ascii=False)},
    )

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(DOCX_MIME)
    assert response.headers["x-annotation-count"] == "1"
    assert response.headers["x-skipped-count"] == "0"
    annotated = Document(io.BytesIO(response.content))
    assert len(annotated.comments) == 1
    assert "f-api-annotation" in next(iter(annotated.comments)).text


def test_api_rejects_invalid_count_and_unsupported_type_with_schema_error():
    client = TestClient(create_app())
    content = _docx_bytes("单文件不足以执行两两比对。")

    count_response = client.post(
        "/v1/text-compare",
        files=[("files", ("single.docx", content, DOCX_MIME))],
    )
    assert count_response.status_code == 400
    assert count_response.json()["error"]["code"] == "invalid_file_count"

    type_response = client.post(
        "/v1/parse",
        files={"file": ("source.txt", b"plain text", "text/plain")},
    )
    assert type_response.status_code == 415
    schema = json.loads((ROOT / "schemas" / "api_error.schema.json").read_text(encoding="utf-8"))
    validate(type_response.json(), schema)


def test_api_enforces_upload_size_limit():
    client = TestClient(create_app(max_upload_bytes=100))
    content = _docx_bytes("此文件必然超过一百字节。")

    response = client.post(
        "/v1/parse",
        files={"file": ("large.docx", content, DOCX_MIME)},
    )

    assert response.status_code == 413
    assert response.json()["error"]["code"] == "file_too_large"


def test_annotate_rejects_malformed_findings_json():
    client = TestClient(create_app())
    content = _docx_bytes("批注输入校验。")

    response = client.post(
        "/v1/annotate-docx",
        files={"file": ("source.docx", content, DOCX_MIME)},
        data={"findings_json": "{not-json"},
    )

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "invalid_findings_json"
