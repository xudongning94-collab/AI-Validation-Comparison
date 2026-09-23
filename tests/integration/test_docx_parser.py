from pathlib import Path

from docx import Document

from bid_compare_agent.parser.docx_parser import parse_docx


def test_parse_docx(tmp_path: Path):
    p = tmp_path / "sample.docx"
    doc = Document()
    doc.add_heading("技术方案", level=1)
    para = doc.add_paragraph("这是第一段方案内容。")
    para.runs[0].bold = True
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).text = "并发"
    table.cell(1, 1).text = "100"
    doc.save(p)

    ir = parse_docx(p)
    assert ir.file_type == "docx"
    assert len(ir.paragraphs) >= 2
    assert ir.paragraphs[0].heading_level == 1
    assert len(ir.tables) == 1
    assert ir.tables[0].rows[1][1] == "100"
    assert ir.sha256
