from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

from docx import Document
from jsonschema import validate

from bid_compare_agent.annotate import annotate_docx
from bid_compare_agent.models.document_ir import SourceLocator
from bid_compare_agent.models.finding import Finding
from bid_compare_agent.parser.docx_parser import parse_docx


ROOT = Path(__file__).resolve().parents[2]


def _source_and_finding(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("需要批注的正文段落。")
    document.save(source)
    parsed = parse_docx(source)
    finding = Finding(
        finding_id="f-schema",
        type="format_issue",
        severity="medium",
        document_id=parsed.document_id,
        source_locator=SourceLocator(kind="docx_paragraph", paragraph_index=0),
        summary="格式需要人工检查",
        score=0.7,
    )
    return source, finding


def test_annotation_result_matches_schema(tmp_path):
    source, finding = _source_and_finding(tmp_path)
    result = annotate_docx(source, tmp_path / "annotated.docx", [finding])
    schema = json.loads((ROOT / "schemas" / "annotation.schema.json").read_text(encoding="utf-8"))

    validate(result.to_dict(), schema)


def test_annotation_cli_creates_docx_and_json_report(tmp_path):
    source, finding = _source_and_finding(tmp_path)
    findings_path = tmp_path / "findings.json"
    output_path = tmp_path / "annotated.docx"
    report_path = tmp_path / "annotation.json"
    findings_path.write_text(
        json.dumps({"findings": [finding.to_dict()]}, ensure_ascii=False),
        encoding="utf-8",
    )

    completed = subprocess.run(
        [
            sys.executable,
            str(ROOT / "scripts" / "annotate_docx.py"),
            str(source),
            str(findings_path),
            "-o",
            str(output_path),
            "--report",
            str(report_path),
        ],
        cwd=ROOT,
        capture_output=True,
        text=True,
        check=False,
    )

    assert completed.returncode == 0, completed.stderr
    assert "OK annotated=1 skipped=0" in completed.stdout
    assert output_path.exists()
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert report["summary"]["annotated_comments"] == 1
    assert len(Document(output_path).comments) == 1
