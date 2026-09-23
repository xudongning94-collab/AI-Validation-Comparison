from pathlib import Path

from docx import Document
from jsonschema import validate
from PIL import Image, ImageDraw

from bid_compare_agent.compare import ImageCompareConfig, compare_image_pair
from bid_compare_agent.parser.docx_parser import parse_docx


def _make_pattern(path: Path, size: int, variant: int = 0) -> None:
    image = Image.new("RGB", (size, size), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((size // 8, size // 8, size * 7 // 8, size * 7 // 8), outline="black", width=max(1, size // 16))
    if variant == 0:
        draw.line((0, 0, size - 1, size - 1), fill="blue", width=max(1, size // 16))
    else:
        draw.ellipse((size // 4, size // 4, size * 3 // 4, size * 3 // 4), fill="red")
    image.save(path)


def _make_doc(path: Path, image_path: Path) -> None:
    doc = Document()
    doc.add_heading("系统架构", level=1)
    p = doc.add_paragraph("整体架构如下图所示：")
    p.add_run().add_picture(str(image_path))
    doc.save(path)


def test_resized_same_visual_image_is_detected(tmp_path: Path):
    image_a = tmp_path / "a.png"
    image_b = tmp_path / "b.png"
    _make_pattern(image_a, 64, variant=0)
    _make_pattern(image_b, 128, variant=0)

    doc_a_path = tmp_path / "a.docx"
    doc_b_path = tmp_path / "b.docx"
    _make_doc(doc_a_path, image_a)
    _make_doc(doc_b_path, image_b)

    doc_a = parse_docx(doc_a_path)
    doc_b = parse_docx(doc_b_path)
    findings, summary = compare_image_pair(doc_a, doc_b, ImageCompareConfig())

    assert len(doc_a.images) == 1
    assert len(doc_b.images) == 1
    assert doc_a.images[0].phash
    assert doc_b.images[0].phash
    assert findings
    assert summary.exact_pairs + summary.high_similar_pairs >= 1
    assert summary.document_a.repeated_rate == 1.0
    assert findings[0].source_locator.paragraph_index is not None


def test_different_image_is_not_detected_as_duplicate(tmp_path: Path):
    image_a = tmp_path / "a.png"
    image_b = tmp_path / "b.png"
    _make_pattern(image_a, 96, variant=0)
    _make_pattern(image_b, 96, variant=1)
    doc_a_path = tmp_path / "a.docx"
    doc_b_path = tmp_path / "b.docx"
    _make_doc(doc_a_path, image_a)
    _make_doc(doc_b_path, image_b)

    findings, summary = compare_image_pair(parse_docx(doc_a_path), parse_docx(doc_b_path))
    assert findings == []
    assert summary.max_similarity == 0.0


def test_image_compare_result_matches_schema(tmp_path: Path):
    from bid_compare_agent.compare import compare_image_set

    image_a = tmp_path / "a.png"
    image_b = tmp_path / "b.png"
    _make_pattern(image_a, 64, variant=0)
    _make_pattern(image_b, 128, variant=0)
    doc_a_path = tmp_path / "a.docx"
    doc_b_path = tmp_path / "b.docx"
    _make_doc(doc_a_path, image_a)
    _make_doc(doc_b_path, image_b)
    result = compare_image_set([parse_docx(doc_a_path), parse_docx(doc_b_path)])

    import json
    schema = json.loads((Path(__file__).resolve().parents[2] / "schemas" / "image_compare.schema.json").read_text(encoding="utf-8"))
    validate(result.to_dict(), schema)
