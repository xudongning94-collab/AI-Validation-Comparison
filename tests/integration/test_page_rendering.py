from __future__ import annotations

import hashlib
import json
import subprocess
import sys
from pathlib import Path
from types import SimpleNamespace

import jsonschema
import pymupdf
from docx import Document
from PIL import Image

from bid_compare_agent.vision import page_rendering


ROOT = Path(__file__).resolve().parents[2]


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _write_pdf(path: Path, page_count: int = 2) -> None:
    document = pymupdf.open()
    for index in range(page_count):
        page = document.new_page(width=320, height=480)
        page.insert_text((40, 60), f"Signed page {index + 1}")
    document.save(path)
    document.close()


def _write_matching_docx_and_pdf(docx_path: Path, pdf_path: Path) -> None:
    sections = [
        (f"Canonical tender section {index} " * 24).strip()
        for index in range(1, 3)
    ]
    document = Document()
    for section in sections:
        document.add_paragraph(section)
    document.save(docx_path)

    pdf = pymupdf.open()
    for section in sections:
        page = pdf.new_page(width=595, height=842)
        page.insert_textbox(
            pymupdf.Rect(50, 50, 545, 792),
            section,
            fontsize=10,
        )
    pdf.save(pdf_path)
    pdf.close()


def _schema() -> dict:
    return json.loads((ROOT / "schemas" / "page_render.schema.json").read_text(encoding="utf-8"))


def test_pdf_pages_render_to_hash_named_images_without_mutating_source(tmp_path: Path) -> None:
    source = tmp_path / "private tender name.pdf"
    output = tmp_path / "rendered"
    _write_pdf(source)
    original_hash = _sha256(source)

    result = page_rendering.render_document_pages(source, output, dpi=144)

    jsonschema.validate(result, _schema())
    assert result["status"] == "ok"
    assert result["source_type"] == "pdf"
    assert result["source_sha256"] == original_hash == _sha256(source)
    assert result["page_count"] == 2
    assert result["renderer"]["backend"] == "pymupdf"
    assert "private tender name" not in json.dumps(result, ensure_ascii=False)
    for page in result["pages"]:
        image_path = output / page["image"]
        assert image_path.is_file()
        assert page["image"].startswith(f"{original_hash[:16]}-page-")
        assert page["image_sha256"] == _sha256(image_path)
        with Image.open(image_path) as image:
            assert image.mode == "RGB"
            assert image.size == (page["width"], page["height"])

    repeated = page_rendering.render_document_pages(source, output, dpi=144)
    assert repeated == result

    cli_output = tmp_path / "cli-rendered"
    completed = subprocess.run(
        [
            sys.executable,
            str(ROOT / "scripts" / "render_document_pages.py"),
            str(source),
            "--output-dir",
            str(cli_output),
        ],
        cwd=ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )
    assert completed.returncode == 0, completed.stderr
    cli_result = json.loads((cli_output / "page-render.json").read_text(encoding="utf-8"))
    jsonschema.validate(cli_result, _schema())
    assert cli_result["status"] == "ok"
    assert str(source) not in json.dumps(cli_result, ensure_ascii=False)


def test_docx_without_renderer_is_explicitly_unavailable_and_writes_no_pages(tmp_path: Path) -> None:
    source = tmp_path / "private response.docx"
    output = tmp_path / "rendered"
    document = Document()
    document.add_paragraph("Signature area")
    document.save(source)
    original_hash = _sha256(source)

    result = page_rendering.render_document_pages(
        source,
        output,
        docx_renderer=tmp_path / "missing-soffice.exe",
    )

    jsonschema.validate(result, _schema())
    assert result["status"] == "unavailable"
    assert result["error"]["code"] == "docx_renderer_unavailable"
    assert result["source_sha256"] == original_hash == _sha256(source)
    assert result["pages"] == []
    assert list(output.glob("*.png")) == []
    assert str(source) not in json.dumps(result, ensure_ascii=False)


def test_docx_renderer_discovers_standard_windows_installation(monkeypatch, tmp_path: Path) -> None:
    program_files = tmp_path / "Program Files"
    renderer = program_files / "LibreOffice" / "program" / "soffice.exe"
    renderer.parent.mkdir(parents=True)
    renderer.write_bytes(b"not executed")
    monkeypatch.setenv("PROGRAMFILES", str(program_files))
    monkeypatch.delenv("PROGRAMFILES(X86)", raising=False)
    monkeypatch.setattr(page_rendering.shutil, "which", lambda _name: None)

    assert page_rendering._find_docx_renderer(None) == renderer.resolve()


def test_page_render_cli_supports_renderer_environment_override() -> None:
    script = (ROOT / "scripts" / "render_document_pages.py").read_text(
        encoding="utf-8"
    )

    assert "docx_renderer_environment_variable" in script
    assert "BID_COMPARE_DOCX_RENDERER" in script


def test_docx_conversion_uses_temporary_pdf_and_preserves_original(monkeypatch, tmp_path: Path) -> None:
    source = tmp_path / "response.docx"
    output = tmp_path / "rendered"
    document = Document()
    document.add_paragraph("Signature area")
    document.save(source)
    original_hash = _sha256(source)

    def fake_convert(_source: Path, work_dir: Path, _renderer: Path, _timeout: int) -> Path:
        converted = work_dir / "converted.pdf"
        _write_pdf(converted, page_count=1)
        return converted

    monkeypatch.setattr(page_rendering, "_convert_docx_to_pdf", fake_convert)
    fake_renderer = tmp_path / "fake-soffice.exe"
    fake_renderer.write_bytes(b"not executed")
    result = page_rendering.render_document_pages(
        source,
        output,
        docx_renderer=fake_renderer,
    )

    jsonschema.validate(result, _schema())
    assert result["status"] == "ok"
    assert result["source_type"] == "docx"
    assert result["source_sha256"] == original_hash == _sha256(source)
    assert result["page_count"] == 1
    assert result["renderer"]["backend"] == "libreoffice+pymupdf"
    assert not any(path.suffix == ".pdf" for path in output.iterdir())


def test_docx_renderer_output_uses_explicit_safe_decoding(monkeypatch, tmp_path: Path) -> None:
    source = tmp_path / "response.docx"
    source.write_bytes(b"docx-placeholder")
    renderer = tmp_path / "soffice.exe"
    renderer.write_bytes(b"renderer-placeholder")

    def fake_run(command, **kwargs):
        assert kwargs["text"] is True
        assert kwargs["encoding"] == "utf-8"
        assert kwargs["errors"] == "replace"
        converted = tmp_path / "response.pdf"
        _write_pdf(converted, page_count=1)
        return SimpleNamespace(returncode=0)

    monkeypatch.setattr(page_rendering.subprocess, "run", fake_run)

    converted = page_rendering._convert_docx_to_pdf(
        source,
        tmp_path,
        renderer,
        timeout_seconds=30,
    )

    assert converted == tmp_path / "response.pdf"


def test_docx_uses_verified_canonical_pdf_without_libreoffice(tmp_path: Path) -> None:
    source = tmp_path / "private response.docx"
    canonical_pdf = tmp_path / "trusted export.pdf"
    output = tmp_path / "rendered"
    _write_matching_docx_and_pdf(source, canonical_pdf)
    source_hash = _sha256(source)
    canonical_hash = _sha256(canonical_pdf)

    result = page_rendering.render_document_pages(
        source,
        output,
        docx_renderer=tmp_path / "missing-soffice.exe",
        canonical_pdf=canonical_pdf,
        canonical_pdf_minimum_text_match=0.7,
    )

    jsonschema.validate(result, _schema())
    assert result["status"] == "ok"
    assert result["source_type"] == "docx"
    assert result["source_sha256"] == source_hash == _sha256(source)
    assert result["page_count"] == 2
    assert result["renderer"]["backend"] == "canonical-pdf+pymupdf"
    assert result["canonical_pdf"]["source_sha256"] == canonical_hash
    assert result["canonical_pdf"]["page_count"] == 2
    assert result["canonical_pdf"]["integrity"]["source_unchanged"] is True
    assert result["canonical_pdf"]["content_match"]["matched"] is True
    assert result["canonical_pdf"]["content_match"]["minimum_bidirectional_match"] == 0.7
    assert result["warnings"] == [
        {
            "code": "canonical_pdf_used",
            "message": "Verified same-source PDF supplied the canonical page layout.",
        }
    ]
    assert result["pages"][0]["image"].startswith(canonical_hash[:16])
    assert _sha256(canonical_pdf) == canonical_hash
    assert str(source) not in json.dumps(result, ensure_ascii=False)
    assert str(canonical_pdf) not in json.dumps(result, ensure_ascii=False)

    cli_output = tmp_path / "cli-rendered"
    completed = subprocess.run(
        [
            sys.executable,
            str(ROOT / "scripts" / "render_document_pages.py"),
            str(source),
            "--canonical-pdf",
            str(canonical_pdf),
            "--output-dir",
            str(cli_output),
        ],
        cwd=ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )
    assert completed.returncode == 0, completed.stderr
    cli_result = json.loads((cli_output / "page-render.json").read_text(encoding="utf-8"))
    jsonschema.validate(cli_result, _schema())
    assert cli_result["renderer"]["backend"] == "canonical-pdf+pymupdf"
    assert cli_result["canonical_pdf"]["content_match"]["matched"] is True


def test_docx_rejects_unrelated_canonical_pdf_without_publishing_pages(tmp_path: Path) -> None:
    source = tmp_path / "private response.docx"
    canonical_pdf = tmp_path / "unrelated export.pdf"
    output = tmp_path / "rendered"
    document = Document()
    document.add_paragraph("Expected tender response content " * 30)
    document.save(source)
    pdf = pymupdf.open()
    page = pdf.new_page(width=595, height=842)
    page.insert_textbox(
        pymupdf.Rect(50, 50, 545, 792),
        "Completely unrelated reference material " * 30,
        fontsize=10,
    )
    pdf.save(canonical_pdf)
    pdf.close()

    result = page_rendering.render_document_pages(
        source,
        output,
        canonical_pdf=canonical_pdf,
        canonical_pdf_minimum_text_match=0.7,
    )

    jsonschema.validate(result, _schema())
    assert result["status"] == "failed"
    assert result["error"]["code"] == "canonical_pdf_content_mismatch"
    assert result["pages"] == []
    assert result["canonical_pdf"]["content_match"]["matched"] is False
    assert result["warnings"] == []
    assert list(output.glob("*.png")) == []
