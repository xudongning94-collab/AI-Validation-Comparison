from pathlib import Path

from docx import Document
from docx.shared import Pt

from bid_compare_agent.check import check_document_format
from bid_compare_agent.parser import parse_document


def test_format_check_finds_outlier(tmp_path: Path):
    path = tmp_path / "mixed.docx"
    doc = Document()
    for i in range(4):
        p = doc.add_paragraph(f"这是正文标准段落{i}，用于验证文档格式检测能力能够识别主流样式。")
        p.runs[0].font.name = "Arial"
        p.runs[0].font.size = Pt(12)
    p = doc.add_paragraph("这是格式异常段落，用于验证字体字号离群检测逻辑。")
    p.runs[0].font.name = "Courier New"
    p.runs[0].font.size = Pt(18)
    doc.save(path)

    ir = parse_document(path)
    result = check_document_format(ir)
    assert any(f.metadata.get("rule") == "body-style-outlier" for f in result.findings)
