from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from bid_compare_agent.models.document_ir import (
    DocumentIR,
    FormatMeta,
    ImageIR,
    ParagraphIR,
    RunStyle,
    SourceLocator,
    TableIR,
)
from bid_compare_agent.utils.io import bytes_sha256, file_sha256
from bid_compare_agent.utils.image import extract_image_features
from bid_compare_agent.utils.text import normalize_text


_ALIGNMENT = {
    WD_ALIGN_PARAGRAPH.LEFT: "left",
    WD_ALIGN_PARAGRAPH.CENTER: "center",
    WD_ALIGN_PARAGRAPH.RIGHT: "right",
    WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
}


def _heading_level(style_name: str | None) -> int | None:
    if not style_name:
        return None
    m = re.search(r"(?:Heading|标题)\s*([1-9])", style_name, re.IGNORECASE)
    return int(m.group(1)) if m else None


def _cm(length) -> float | None:
    if length is None:
        return None
    return round(float(length.cm), 3)


def parse_docx(path: str | Path, schema_version: str = "1.0.0") -> DocumentIR:
    path = Path(path)
    doc = Document(path)
    paragraphs: list[ParagraphIR] = []
    tables: list[TableIR] = []
    images: list[ImageIR] = []
    warnings: list[str] = []

    for idx, p in enumerate(doc.paragraphs):
        text = normalize_text(p.text)
        if not text:
            continue
        runs = []
        for r in p.runs:
            if not r.text:
                continue
            size = float(r.font.size.pt) if r.font.size else None
            runs.append(
                RunStyle(
                    text=r.text,
                    bold=r.bold,
                    italic=r.italic,
                    underline=bool(r.underline) if r.underline is not None else None,
                    font_name=r.font.name,
                    font_size_pt=size,
                )
            )
        spacing = p.paragraph_format.line_spacing
        if hasattr(spacing, "pt"):
            spacing_value = f"{round(float(spacing.pt), 2)}pt"
        elif isinstance(spacing, (int, float)):
            spacing_value = float(spacing)
        else:
            spacing_value = None
        paragraphs.append(
            ParagraphIR(
                id=f"p-{len(paragraphs)+1:06d}",
                text=text,
                order=len(paragraphs),
                style_name=p.style.name if p.style else None,
                heading_level=_heading_level(p.style.name if p.style else None),
                alignment=_ALIGNMENT.get(p.alignment),
                line_spacing=spacing_value,
                runs=runs,
                source_locator=SourceLocator(kind="docx_paragraph", paragraph_index=idx, index=idx),
            )
        )

    for t_idx, table in enumerate(doc.tables):
        rows = [[normalize_text(cell.text) for cell in row.cells] for row in table.rows]
        tables.append(
            TableIR(
                id=f"t-{t_idx+1:05d}",
                order=t_idx,
                rows=rows,
                source_locator=SourceLocator(kind="docx_table", table_index=t_idx, index=t_idx),
            )
        )

    # 建立图片 relationship -> 正文段落索引映射，后续图片重复问题可以尽量回到原文位置。
    image_rel_to_paragraph: dict[str, int] = {}
    embed_attr = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
    for paragraph_index, paragraph in enumerate(doc.paragraphs):
        try:
            for blip in paragraph._p.xpath(".//a:blip"):
                rel_id = blip.get(embed_attr)
                if rel_id and rel_id not in image_rel_to_paragraph:
                    image_rel_to_paragraph[rel_id] = paragraph_index
        except Exception:
            # 个别复杂 OOXML 结构无法解析时仍保留 relationship 级定位。
            continue

    image_order = 0
    for rel_id, rel in doc.part.rels.items():
        if "image" not in rel.reltype:
            continue
        try:
            blob = rel.target_part.blob
            filename = Path(str(rel.target_ref)).name
            content_type = getattr(rel.target_part, "content_type", None)
            features = extract_image_features(blob)
            paragraph_index = image_rel_to_paragraph.get(rel_id)
            images.append(
                ImageIR(
                    id=f"img-{image_order+1:05d}",
                    order=image_order,
                    filename=filename,
                    content_type=content_type,
                    width_px=features.width_px,
                    height_px=features.height_px,
                    sha256=bytes_sha256(blob),
                    phash=features.phash,
                    mean_rgb=features.mean_rgb,
                    source_locator=SourceLocator(
                        kind="docx_relationship",
                        index=image_order,
                        paragraph_index=paragraph_index,
                        part=str(rel.target_ref),
                    ),
                )
            )
            image_order += 1
        except Exception as exc:
            warnings.append(f"图片提取失败: {filename if 'filename' in locals() else rel_id}: {exc}")

    fm = FormatMeta(section_count=len(doc.sections))
    if doc.sections:
        s = doc.sections[0]
        fm.page_width_cm = _cm(s.page_width)
        fm.page_height_cm = _cm(s.page_height)
        fm.margin_top_cm = _cm(s.top_margin)
        fm.margin_bottom_cm = _cm(s.bottom_margin)
        fm.margin_left_cm = _cm(s.left_margin)
        fm.margin_right_cm = _cm(s.right_margin)
        for sec in doc.sections:
            h = normalize_text("\n".join(p.text for p in sec.header.paragraphs))
            f = normalize_text("\n".join(p.text for p in sec.footer.paragraphs))
            if h and h not in fm.header_texts:
                fm.header_texts.append(h)
            if f and f not in fm.footer_texts:
                fm.footer_texts.append(f)

    props = doc.core_properties
    metadata = {
        "title": props.title or None,
        "subject": props.subject or None,
        "author": props.author or None,
        "last_modified_by": props.last_modified_by or None,
        "created": props.created.isoformat() if props.created else None,
        "modified": props.modified.isoformat() if props.modified else None,
    }

    file_hash = file_sha256(path)

    return DocumentIR(
        schema_version=schema_version,
        document_id=f"doc-{file_hash[:24]}",
        filename=path.name,
        file_type="docx",
        file_size_bytes=path.stat().st_size,
        sha256=file_hash,
        page_count=None,
        paragraphs=paragraphs,
        tables=tables,
        images=images,
        format_meta=fm,
        warnings=warnings,
        metadata=metadata,
    )
