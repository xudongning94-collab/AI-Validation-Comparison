from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.text.paragraph import Paragraph

from bid_compare_agent.models.annotation import AnnotationRecord, AnnotationSkip, DocxAnnotationResult
from bid_compare_agent.models.document_ir import SourceLocator
from bid_compare_agent.models.finding import Finding
from bid_compare_agent.parser.docx_parser import parse_docx
from bid_compare_agent.utils.io import file_sha256


_SEVERITY_ORDER = {"info": 0, "low": 1, "medium": 2, "high": 3, "critical": 4}


@dataclass(frozen=True)
class DocxAnnotationConfig:
    author: str = "Bid Compare Agent"
    initials: str = "BCA"
    min_severity: str = "low"
    max_comments: int = 500
    include_evidence: bool = True
    max_evidence_chars: int = 500

    def __post_init__(self) -> None:
        if not self.author.strip():
            raise ValueError("author 不能为空")
        if not self.initials.strip():
            raise ValueError("initials 不能为空")
        if self.min_severity not in _SEVERITY_ORDER:
            raise ValueError(f"不支持的最低严重度: {self.min_severity}")
        if self.max_comments < 1:
            raise ValueError("max_comments 必须大于 0")
        if self.max_evidence_chars < 0:
            raise ValueError("max_evidence_chars 不能小于 0")


def _target_for_document(
    finding: Finding,
    document_id: str,
) -> tuple[str, SourceLocator] | None:
    if finding.document_id == document_id:
        return "source", finding.source_locator
    if finding.peer_document_id == document_id and finding.peer_source_locator is not None:
        return "peer", finding.peer_source_locator
    return None


def _paragraph_for_locator(
    document: DocumentObject,
    locator: SourceLocator,
) -> tuple[Paragraph | None, str | None]:
    if locator.paragraph_index is not None:
        index = locator.paragraph_index
        if not 0 <= index < len(document.paragraphs):
            return None, "paragraph_index_out_of_range"
        return document.paragraphs[index], None

    if locator.table_index is not None:
        table_index = locator.table_index
        if not 0 <= table_index < len(document.tables):
            return None, "table_index_out_of_range"
        table = document.tables[table_index]
        row_index = locator.row_index if locator.row_index is not None else 0
        cell_index = locator.cell_index if locator.cell_index is not None else 0
        if not 0 <= row_index < len(table.rows):
            return None, "row_index_out_of_range"
        if not 0 <= cell_index < len(table.rows[row_index].cells):
            return None, "cell_index_out_of_range"
        cell = table.rows[row_index].cells[cell_index]
        if not cell.paragraphs:
            return None, "table_cell_has_no_paragraph"
        return cell.paragraphs[0], None

    return None, "locator_not_supported"


def _evidence_text(finding: Finding, max_chars: int) -> str:
    if not finding.evidence or max_chars == 0:
        return ""
    serialized = json.dumps(finding.evidence, ensure_ascii=False, sort_keys=True, default=str)
    if len(serialized) <= max_chars:
        return serialized
    return serialized[: max(0, max_chars - 1)] + "…"


def _comment_text(finding: Finding, target_role: str, config: DocxAnnotationConfig) -> str:
    lines = [
        f"[{finding.severity.upper()}] {finding.summary}",
        f"类型: {finding.type}",
        f"Finding ID: {finding.finding_id}",
        f"定位角色: {'主文档' if target_role == 'source' else '对端文档'}",
    ]
    if finding.score is not None:
        lines.append(f"风险/相似度分: {finding.score:.1%}")
    if finding.peer_document_id:
        lines.append(f"关联文档: {finding.peer_document_id}")
    if config.include_evidence:
        evidence = _evidence_text(finding, config.max_evidence_chars)
        if evidence:
            lines.append(f"证据: {evidence}")
    if finding.type == "ai_likelihood":
        lines.append("提示: AI 疑似度仅为辅助信号，不能证明文本由 AI 生成，必须人工复核。")
    return "\n".join(lines)


def annotate_docx(
    source_path: str | Path,
    output_path: str | Path,
    findings: list[Finding],
    config: DocxAnnotationConfig | None = None,
) -> DocxAnnotationResult:
    config = config or DocxAnnotationConfig()
    source_path = Path(source_path)
    output_path = Path(output_path)
    if source_path.suffix.lower() != ".docx":
        raise ValueError("D9 Word 批注当前仅支持 .docx 文件")
    if output_path.suffix.lower() != ".docx":
        raise ValueError("D9 Word 批注输出必须是 .docx 文件")
    if source_path.resolve() == output_path.resolve():
        raise ValueError("输出文件必须与源文件不同，避免覆盖原始投标文件")

    source_ir = parse_docx(source_path)
    document = Document(source_path)
    annotations: list[AnnotationRecord] = []
    skipped: list[AnnotationSkip] = []
    seen: set[tuple[str, str]] = set()

    ordered = sorted(
        findings,
        key=lambda item: (_SEVERITY_ORDER.get(item.severity, -1), item.score or 0.0, item.finding_id),
        reverse=True,
    )
    for finding in ordered:
        target = _target_for_document(finding, source_ir.document_id)
        if target is None:
            skipped.append(
                AnnotationSkip(
                    finding_id=finding.finding_id,
                    finding_type=finding.type,
                    reason="finding_targets_another_document",
                )
            )
            continue
        target_role, locator = target
        dedupe_key = (finding.finding_id, target_role)
        if dedupe_key in seen:
            skipped.append(
                AnnotationSkip(
                    finding_id=finding.finding_id,
                    finding_type=finding.type,
                    reason="duplicate_finding",
                    target_role=target_role,
                    source_locator=locator,
                )
            )
            continue
        seen.add(dedupe_key)

        if _SEVERITY_ORDER.get(finding.severity, -1) < _SEVERITY_ORDER[config.min_severity]:
            skipped.append(
                AnnotationSkip(
                    finding_id=finding.finding_id,
                    finding_type=finding.type,
                    reason="below_min_severity",
                    target_role=target_role,
                    source_locator=locator,
                )
            )
            continue
        if len(annotations) >= config.max_comments:
            skipped.append(
                AnnotationSkip(
                    finding_id=finding.finding_id,
                    finding_type=finding.type,
                    reason="max_comments_reached",
                    target_role=target_role,
                    source_locator=locator,
                )
            )
            continue

        paragraph, reason = _paragraph_for_locator(document, locator)
        if paragraph is None:
            skipped.append(
                AnnotationSkip(
                    finding_id=finding.finding_id,
                    finding_type=finding.type,
                    reason=reason or "paragraph_not_found",
                    target_role=target_role,
                    source_locator=locator,
                )
            )
            continue
        if not paragraph.runs:
            skipped.append(
                AnnotationSkip(
                    finding_id=finding.finding_id,
                    finding_type=finding.type,
                    reason="paragraph_has_no_runs",
                    target_role=target_role,
                    source_locator=locator,
                )
            )
            continue

        comment_text = _comment_text(finding, target_role, config)
        comment = document.add_comment(
            paragraph.runs,
            text=comment_text,
            author=config.author,
            initials=config.initials,
        )
        annotations.append(
            AnnotationRecord(
                finding_id=finding.finding_id,
                finding_type=finding.type,
                severity=finding.severity,
                target_role=target_role,
                source_locator=locator,
                comment_id=comment.comment_id,
                comment_text=comment_text,
            )
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return DocxAnnotationResult(
        schema_version="1.0.0",
        annotation_type="docx_comments",
        source_document_id=source_ir.document_id,
        source_filename=source_path.name,
        output_filename=output_path.name,
        annotations=annotations,
        skipped=skipped,
        summary={
            "input_findings": len(findings),
            "annotated_comments": len(annotations),
            "skipped_findings": len(skipped),
            "skip_reasons": {
                reason: sum(1 for item in skipped if item.reason == reason)
                for reason in sorted({item.reason for item in skipped})
            },
        },
        metadata={
            "source_sha256": source_ir.sha256,
            "output_sha256": file_sha256(output_path),
            "author": config.author,
            "initials": config.initials,
            "min_severity": config.min_severity,
            "max_comments": config.max_comments,
        },
    )
